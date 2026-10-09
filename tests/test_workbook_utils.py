"""Test di scripts/workbook_utils.py: salvare con openpyxl non deve perdere il collegamento a Claude per Excel
e deve lasciare in cache i valori delle formule, come farebbe Excel."""
import shutil
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
import openpyxl  # noqa: E402

import workbook_utils as wu  # noqa: E402

XLSX = ROOT / "data" / "fondi-pensione-covip.xlsx"


class TestSalva(unittest.TestCase):
    def test_conserva_webextensions_e_formule(self):
        with zipfile.ZipFile(XLSX) as z:
            if not any(n.startswith("xl/webextensions/") for n in z.namelist()):
                self.skipTest("il workbook non ha componenti aggiuntivi")
        with tempfile.TemporaryDirectory() as d:
            copia = Path(d) / "copia.xlsx"
            shutil.copy(XLSX, copia)
            wb = openpyxl.load_workbook(copia)
            wu.salva(wb, copia)
            with zipfile.ZipFile(copia) as z:
                nomi = z.namelist()
                self.assertEqual(nomi[0], "[Content_Types].xml")
                self.assertIn("xl/webextensions/webextension1.xml", nomi)
                self.assertIn(wu.REL_TASKPANES, z.read("_rels/.rels").decode())
                ct = z.read("[Content_Types].xml").decode()
                self.assertEqual(ct.count('PartName="/xl/webextensions/taskpanes.xml"'), 1)
            wb2 = openpyxl.load_workbook(copia)
            self.assertTrue(str(wb2["Sheet1"]["H3"].value).startswith("=COUNTIF"))


class TestCacheFormule(unittest.TestCase):
    """riempi_cache() calcola le formule come Excel: chi legge il file con data_only=True trova i valori."""

    def test_formule_del_workbook(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "prova.xlsx"
            wb = openpyxl.Workbook()
            f = wb.active
            f.title = "Sheet1"
            for r, nome in ((3, "Alfa"), (4, "Beta"), (5, "Gamma*")):
                f[f"A{r}"] = nome
                f[f"H{r}"] = f"=COUNTIF(Comparti!$A:$A,$A{r})"
                f[f"I{r}"] = f'=IF($H{r}=0,"",_xlfn.MINIFS(Comparti!$C:$C,Comparti!$A:$A,$A{r}))'
                f[f"J{r}"] = f'=IF($H{r}=0,"",_xlfn.MAXIFS(Comparti!$C:$C,Comparti!$A:$A,$A{r}))'
                f[f"L{r}"] = (f'=IF(COUNTIFS(Comparti!$A:$A,$A{r},Comparti!$B:$B,10)=0,"",'
                              f'_xlfn.MAXIFS(Comparti!$D:$D,Comparti!$A:$A,$A{r},Comparti!$B:$B,10))')
            c = wb.create_sheet("Comparti")
            c.append(["Fondo", "Periodo", "Commissione", "Rendimento"])
            for riga, periodo, comm, rend in ((3, 10, 0.012, 0.03), (3, 5, 0.009, 0.05), (4, 3, 0.02, None)):
                c.append([f"=Sheet1!$A${riga}", periodo, comm, rend])
            v = wb.create_sheet("Varie")
            v["A1"], v["B1"], v["C1"], v["D1"] = 200, 50, "Sì (5 o 10 anni)", "No"
            v["E1"], v["F1"] = "sì", None
            v["A2"] = "=B1/A$1"
            v["A3"] = '=COUNTIF(C1:F1,"Sì*")'
            v["A4"] = '=IF(F1="","",F1/10000)'
            v["A5"] = "=A1/(B1-50)"
            v["A6"] = '=A1&" €"'
            self.assertEqual(wu.salva(wb, p), [])

            wc = openpyxl.load_workbook(p, data_only=True)
            s1, var = wc["Sheet1"], wc["Varie"]
            self.assertEqual([s1[f"H{r}"].value for r in (3, 4, 5)], [2, 1, 0])
            self.assertEqual((s1["I3"].value, s1["J3"].value), (0.009, 0.012))
            self.assertEqual(s1["L3"].value, 0.03)          # solo il comparto a 10 anni
            self.assertIsNone(s1["L4"].value)                # nessun comparto a 10 anni → ""
            self.assertIsNone(s1["I5"].value)                # nessun comparto → ""
            self.assertEqual(wc["Comparti"]["A2"].value, "Alfa")
            self.assertEqual(var["A2"].value, 0.25)
            self.assertEqual(var["A3"].value, 2)             # "Sì (…)" e "sì": jolly e maiuscole come in Excel
            self.assertIsNone(var["A4"].value)
            self.assertEqual(var["A5"].value, "#DIV/0!")
            self.assertEqual(var["A6"].value, "200 €")
            # le formule restano formule
            self.assertEqual(openpyxl.load_workbook(p)["Varie"]["A2"].value, "=B1/A$1")

    def test_funzione_non_supportata(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "prova.xlsx"
            wb = openpyxl.Workbook()
            wb.active["A1"] = "=VLOOKUP(1,B1:C2,2,0)"
            wb.active["A2"] = "=1+1"
            non_calcolate = wu.salva(wb, p)
            self.assertEqual(len(non_calcolate), 1)
            self.assertIn("VLOOKUP", non_calcolate[0])
            ws = openpyxl.load_workbook(p, data_only=True).active
            self.assertIsNone(ws["A1"].value)
            self.assertEqual(ws["A2"].value, 2)

    def test_workbook_reale(self):
        """Tutte le formule del workbook sono calcolabili; dopo salva() l'export non trova cache vuote e i valori
        coincidono con quelli che ricalcola da solo (ricalcola_metriche)."""
        import export_xlsx as ex
        with tempfile.TemporaryDirectory() as d:
            copia = Path(d) / "copia.xlsx"
            shutil.copy(XLSX, copia)
            originale_cache = openpyxl.load_workbook(copia, data_only=True)
            self.assertEqual(wu.salva(openpyxl.load_workbook(copia), copia), [])
            # dove Excel aveva già messo un valore in cache, il calcolo in Python deve dare lo stesso risultato
            nuova_cache = openpyxl.load_workbook(copia, data_only=True)
            formule = openpyxl.load_workbook(copia)
            for ws in formule.worksheets:
                for riga in ws.iter_rows():
                    for c in riga:
                        if isinstance(c.value, str) and c.value.startswith("="):
                            prima = originale_cache[ws.title][c.coordinate].value
                            dopo = nuova_cache[ws.title][c.coordinate].value
                            if prima is not None and prima != "":
                                if isinstance(prima, float):
                                    self.assertAlmostEqual(prima, dopo, places=12, msg=f"{ws.title}!{c.coordinate}")
                                else:
                                    self.assertEqual(prima, dopo, f"{ws.title}!{c.coordinate}")
            rep = ex.Report()
            ex.esporta(copia, rep)
            self.assertEqual(rep.errori, [])
            self.assertFalse([w for w in rep.warning if "cache" in w], rep.warning)


if __name__ == "__main__":
    unittest.main()
