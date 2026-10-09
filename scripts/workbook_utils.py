"""Utilità per modificare il workbook con openpyxl senza perdere ciò che openpyxl non gestisce.

`salva(wb, percorso)`:
1. salva il workbook con openpyxl;
2. **riempie la cache dei valori delle formule** (`riempi_cache`): Excel salva ogni formula insieme al suo ultimo
   risultato, openpyxl invece non sa calcolare e lascia il risultato vuoto. Chi legge il file senza ricalcolare
   (l'export con `data_only=True`, il visualizzatore di VS Code, le anteprime) vedrebbe celle vuote. Qui le formule
   del workbook vengono calcolate in Python (solo le funzioni elencate in `FUNZIONI`) e il risultato viene scritto
   accanto a ogni formula, come farebbe Excel. Excel le ricalcola comunque all'apertura (`fullCalcOnLoad`);
3. reinserisce le parti `xl/webextensions/*` (il collegamento al componente aggiuntivo Claude per Excel), che
   openpyxl scarta, copiandole dal file originale.
"""
from __future__ import annotations

import math
import os
import re
import shutil
import tempfile
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

TIPI_WEBEXT = {
    "/xl/webextensions/taskpanes.xml": "application/vnd.ms-office.webextensiontaskpanes+xml",
    "/xl/webextensions/webextension1.xml": "application/vnd.ms-office.webextension+xml",
}
REL_TASKPANES = "http://schemas.microsoft.com/office/2011/relationships/webextensiontaskpanes"


def salva(wb, percorso: str | Path, originale: str | Path | None = None) -> list[str]:
    """Salva `wb` in `percorso`, riempie la cache delle formule e reinserisce le parti xl/webextensions di
    `originale` (default: `percorso`). Restituisce le celle con formula che non è stato possibile calcolare."""
    percorso = Path(percorso)
    originale = Path(originale) if originale else percorso
    fd, tmp = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    try:
        wb.save(tmp)
        non_calcolate = riempi_cache(Path(tmp))
        if originale.exists():
            reinserisci_webextensions(originale, Path(tmp))
        shutil.move(tmp, percorso)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)
    return non_calcolate


def reinserisci_webextensions(originale: Path, nuovo: Path) -> bool:
    """Copia in `nuovo` le parti xl/webextensions di `originale` con relazioni e content type. True se ne ha copiate."""
    with zipfile.ZipFile(originale) as zo:
        parti = {n: zo.read(n) for n in zo.namelist() if n.startswith("xl/webextensions/")}
    if not parti:
        return False
    with zipfile.ZipFile(nuovo) as zn:
        contenuti = {n: zn.read(n) for n in zn.namelist()}
    contenuti.update(parti)

    rels = contenuti["_rels/.rels"].decode()
    if REL_TASKPANES not in rels:
        nuovo_id = max((int(x) for x in re.findall(r'Id="rId(\d+)"', rels)), default=0) + 1
        rels = rels.replace("</Relationships>", f'<Relationship Id="rId{nuovo_id}" Type="{REL_TASKPANES}" '
                                                 'Target="xl/webextensions/taskpanes.xml"/></Relationships>')
        contenuti["_rels/.rels"] = rels.encode()

    ct = contenuti["[Content_Types].xml"].decode()
    for parte, tipo in TIPI_WEBEXT.items():
        if parte.lstrip("/") in parti and f'PartName="{parte}"' not in ct:
            ct = ct.replace("</Types>", f'<Override PartName="{parte}" ContentType="{tipo}"/></Types>')
    contenuti["[Content_Types].xml"] = ct.encode()

    ordine = ["[Content_Types].xml"] + [n for n in contenuti if n != "[Content_Types].xml"]
    with zipfile.ZipFile(nuovo, "w", zipfile.ZIP_DEFLATED) as z:
        for n in ordine:
            z.writestr(n, contenuti[n])
    return True


# ---------------------------------------------------------------- cache dei valori delle formule

class ErroreExcel(str):
    """Valore di errore di Excel (#DIV/0!, #VALUE!…): in cache si scrive con t="e"."""


class NonSupportata(Exception):
    """Formula con una funzione o una sintassi che questo modulo non sa calcolare (la cache resta vuota)."""


RE_TOKEN = re.compile(r"""\s*(?:
      (?P<str>"(?:[^"]|"")*")
    | (?P<rif>(?:(?:'(?:[^']|'')+'|[A-Za-z_][\w.]*)!)?\$?[A-Z]{1,3}\$?\d*(?::\$?[A-Z]{1,3}\$?\d*)?)(?![\w(.])
    | (?P<num>\d+(?:\.\d+)?(?:[eE][+-]?\d+)?)
    | (?P<fun>(?:_xlfn\.)?[A-Z][A-Z0-9.]*)\(
    | (?P<op><>|<=|>=|[-+*/&=<>(),])
    )""", re.X)
RE_CELLA = re.compile(r"\$?([A-Z]{1,3})\$?(\d*)")
FUNZIONI = ("IF", "COUNTIF", "COUNTIFS", "MINIFS", "MAXIFS")   # le sole usate nel workbook (CLAUDE.md §3)


def _colonna(lettere: str) -> int:
    n = 0
    for ch in lettere:
        n = n * 26 + ord(ch) - 64
    return n


class _Intervallo:
    """Riferimento a una cella o a un intervallo (anche colonne intere, es. Comparti!$A:$A)."""

    def __init__(self, foglio: str, c1: int, r1: int | None, c2: int, r2: int | None):
        self.foglio, self.c1, self.r1, self.c2, self.r2 = foglio, c1, r1, c2, r2


def _numero(v):
    """Conversione a numero come nelle operazioni aritmetiche di Excel (cella vuota = 0)."""
    if v is None:
        return 0
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, str):
        try:
            return float(v.replace(",", ".")) if v.strip() else ErroreExcel("#VALUE!")
        except ValueError:
            return ErroreExcel("#VALUE!")
    raise NonSupportata(f"valore di tipo {type(v).__name__}")


def _confronta(a, b) -> int:
    """Confronto tra valori con le regole di Excel: numeri < testo < logici, testo senza maiuscole/minuscole,
    cella vuota uguale a 0, "" o FALSO a seconda dell'altro operando."""
    def norm(v, altro):
        if v is None:
            v = "" if isinstance(altro, str) else False if isinstance(altro, bool) else 0
        if isinstance(v, bool):
            return (2, v)
        if isinstance(v, (int, float)):
            return (0, v)
        if isinstance(v, str):
            return (1, v.casefold())
        raise NonSupportata(f"confronto con un valore di tipo {type(v).__name__}")
    x, y = norm(a, b), norm(b, a)
    return (x > y) - (x < y)


def _criterio(c):
    """Funzione cella → bool equivalente al criterio di COUNTIF/MINIFS… (operatori, caratteri jolly * e ?)."""
    if isinstance(c, ErroreExcel):
        return lambda v: False
    if c is None:
        raise NonSupportata("criterio vuoto")
    if isinstance(c, bool):
        return lambda v: isinstance(v, bool) and v == c
    if isinstance(c, (int, float)):
        return lambda v: isinstance(v, (int, float)) and not isinstance(v, bool) and v == c
    op, val = re.match(r"^(<=|>=|<>|<|>|=)?(.*)$", str(c), re.S).groups()
    op = op or "="
    try:
        num = float(val.replace(",", ".")) if val.strip() else None
    except ValueError:
        num = None
    confronti = {"=": lambda d: d == 0, "<>": lambda d: d != 0, "<": lambda d: d < 0, ">": lambda d: d > 0,
                 "<=": lambda d: d <= 0, ">=": lambda d: d >= 0}
    if num is not None:
        def ok(v):
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                return confronti[op]((v > num) - (v < num))
            return op == "<>"
        return ok
    if op in ("=", "<>"):
        if val == "":
            vuota = lambda v: v is None or v == ""  # noqa: E731
            return vuota if op == "=" else (lambda v: not vuota(v))
        # caratteri jolly: * (qualsiasi sequenza), ? (un carattere), ~ li rende letterali
        regex = "".join(re.escape(p[1]) if p.startswith("~") else ".*" if p == "*" else "." if p == "?" else re.escape(p)
                        for p in re.split(r"(~[*?~]|[*?])", val) if p)
        rx = re.compile(f"^{regex}$", re.I | re.S)
        uguale = lambda v: isinstance(v, str) and rx.match(v) is not None  # noqa: E731
        return uguale if op == "=" else (lambda v: not uguale(v))
    return lambda v: isinstance(v, str) and confronti[op](_confronta(v, val))


class Calcolatore:
    """Calcola le formule di un workbook openpyxl (aperto senza data_only), con memoria dei risultati."""

    def __init__(self, wb):
        self.wb = wb
        self.memo: dict[tuple[str, int, int], object] = {}
        self.in_corso: set[tuple[str, int, int]] = set()

    # -- celle e intervalli
    def valore(self, foglio: str, riga: int, col: int):
        chiave = (foglio, riga, col)
        if chiave in self.memo:
            return self.memo[chiave]
        if foglio not in self.wb.sheetnames:
            raise NonSupportata(f"foglio {foglio!r} inesistente")
        v = self.wb[foglio].cell(riga, col).value
        if isinstance(v, str) and v.startswith("="):
            if chiave in self.in_corso:
                raise NonSupportata("riferimento circolare")
            self.in_corso.add(chiave)
            try:
                v = self.formula(v[1:], foglio)
            finally:
                self.in_corso.discard(chiave)
        elif v is not None and not isinstance(v, (str, int, float, bool)):
            v = str(v) if not hasattr(v, "isoformat") else v   # date: non servono alle formule attuali
        self.memo[chiave] = v
        return v

    def celle(self, iv: _Intervallo) -> list:
        ws = self.wb[iv.foglio]
        r1 = iv.r1 or 1
        r2 = iv.r2 or ws.max_row
        return [self.valore(iv.foglio, r, c) for r in range(r1, r2 + 1) for c in range(iv.c1, iv.c2 + 1)]

    def scalare(self, v):
        if isinstance(v, _Intervallo):
            if v.c1 == v.c2 and v.r1 is not None and v.r1 == v.r2:
                return self.valore(v.foglio, v.r1, v.c1)
            raise NonSupportata("intervallo usato come valore singolo")
        return v

    # -- analisi della formula
    def formula(self, testo: str, foglio: str):
        tok = []
        pos = 0
        while pos < len(testo):
            m = RE_TOKEN.match(testo, pos)
            if not m or m.end() == pos:
                if testo[pos:].strip() == "":
                    break
                raise NonSupportata(f"sintassi non riconosciuta: {testo[pos:pos + 20]!r}")
            tok.append((m.lastgroup, m.group(m.lastgroup)))
            pos = m.end()
        # le celle citate possono contenere a loro volta formule: lo stato dell'analisi va salvato e ripristinato
        stato = getattr(self, "tok", None), getattr(self, "i", 0), getattr(self, "foglio", None)
        self.tok, self.i, self.foglio = tok, 0, foglio
        try:
            v = self.espressione()
            if self.i != len(self.tok):
                raise NonSupportata(f"token in eccesso: {self.tok[self.i:]}")
            return self.scalare(v)
        finally:
            self.tok, self.i, self.foglio = stato

    def guarda(self):
        return self.tok[self.i] if self.i < len(self.tok) else (None, None)

    def prendi(self, atteso=None):
        t = self.guarda()
        if atteso is not None and t[1] != atteso:
            raise NonSupportata(f"atteso {atteso!r}, trovato {t[1]!r}")
        self.i += 1
        return t

    def espressione(self):
        a = self.concatena()
        while self.guarda() in (("op", "="), ("op", "<>"), ("op", "<"), ("op", ">"), ("op", "<="), ("op", ">=")):
            op = self.prendi()[1]
            b = self.concatena()
            a, b = self.scalare(a), self.scalare(b)
            for x in (a, b):
                if isinstance(x, ErroreExcel):
                    return x
            d = _confronta(a, b)
            a = {"=": d == 0, "<>": d != 0, "<": d < 0, ">": d > 0, "<=": d <= 0, ">=": d >= 0}[op]
        return a

    def concatena(self):
        a = self.somma()
        while self.guarda() == ("op", "&"):
            self.prendi()
            b = self.scalare(self.somma())
            a = self.scalare(a)
            if isinstance(a, ErroreExcel) or isinstance(b, ErroreExcel):
                a = a if isinstance(a, ErroreExcel) else b
                continue
            a = "".join("" if x is None else ("VERO" if x else "FALSO") if isinstance(x, bool)
                        else (str(int(x)) if isinstance(x, float) and x.is_integer() else str(x)) for x in (a, b))
        return a

    def _aritmetica(self, a, b, op):
        a, b = _numero(self.scalare(a)), _numero(self.scalare(b))
        for x in (a, b):
            if isinstance(x, ErroreExcel):
                return x
        if op == "/":
            if b == 0:
                return ErroreExcel("#DIV/0!")
            r = a / b
        else:
            r = {"+": a + b, "-": a - b, "*": a * b}[op]
        return r if not (isinstance(r, float) and (math.isinf(r) or math.isnan(r))) else ErroreExcel("#NUM!")

    def somma(self):
        a = self.prodotto()
        while self.guarda() in (("op", "+"), ("op", "-")):
            op = self.prendi()[1]
            a = self._aritmetica(a, self.prodotto(), op)
        return a

    def prodotto(self):
        a = self.unario()
        while self.guarda() in (("op", "*"), ("op", "/")):
            op = self.prendi()[1]
            a = self._aritmetica(a, self.unario(), op)
        return a

    def unario(self):
        if self.guarda() == ("op", "-"):
            self.prendi()
            return self._aritmetica(0, self.unario(), "-")
        if self.guarda() == ("op", "+"):
            self.prendi()
        return self.primario()

    def primario(self):
        tipo, testo = self.prendi()
        if tipo == "num":
            return float(testo) if any(ch in testo for ch in ".eE") else int(testo)
        if tipo == "str":
            return testo[1:-1].replace('""', '"')
        if tipo == "rif":
            return self.riferimento(testo)
        if tipo == "op" and testo == "(":
            v = self.espressione()
            self.prendi(")")
            return v
        if tipo == "fun":
            nome = testo.removeprefix("_xlfn.").upper()
            args = []
            if self.guarda() != ("op", ")"):
                args.append(self.espressione())
                while self.guarda() == ("op", ","):
                    self.prendi()
                    args.append(self.espressione())
            self.prendi(")")
            return self.funzione(nome, args)
        raise NonSupportata(f"token inatteso {testo!r}")

    def riferimento(self, testo: str) -> _Intervallo:
        foglio = self.foglio
        if "!" in testo:
            foglio, testo = testo.rsplit("!", 1)
            if foglio.startswith("'"):
                foglio = foglio[1:-1].replace("''", "'")
        parti = [RE_CELLA.fullmatch(p) for p in testo.split(":")]
        if any(p is None for p in parti):
            raise NonSupportata(f"riferimento non riconosciuto: {testo!r}")
        estremi = [(_colonna(p.group(1)), int(p.group(2)) if p.group(2) else None) for p in parti]
        (c1, r1), (c2, r2) = estremi * 2 if len(estremi) == 1 else estremi
        if (r1 is None) != (r2 is None) or (len(estremi) == 1 and r1 is None):
            raise NonSupportata(f"riferimento misto: {testo!r}")
        return _Intervallo(foglio, min(c1, c2), min(r1, r2) if r1 else None, max(c1, c2), max(r1, r2) if r2 else None)

    # -- funzioni
    def funzione(self, nome: str, args: list):
        if nome not in FUNZIONI:
            raise NonSupportata(f"funzione {nome} non supportata")
        if nome == "IF":
            if not 2 <= len(args) <= 3:
                raise NonSupportata("IF con un numero di argomenti inatteso")
            cond = self.scalare(args[0])
            if isinstance(cond, ErroreExcel):
                return cond
            if isinstance(cond, str):
                return ErroreExcel("#VALUE!")
            scelto = args[1] if (bool(cond) if cond is not None else False) else (args[2] if len(args) == 3 else False)
            return self.scalare(scelto)
        if nome == "COUNTIF":
            nome, args = "COUNTIFS", args
        bersaglio = None
        if nome in ("MINIFS", "MAXIFS"):
            bersaglio, args = args[0], args[1:]
        if not args or len(args) % 2:
            raise NonSupportata(f"{nome} con un numero di argomenti inatteso")
        coppie = []
        for rng, crit in zip(args[::2], args[1::2]):
            if not isinstance(rng, _Intervallo):
                raise NonSupportata(f"{nome}: il primo argomento di ogni coppia deve essere un intervallo")
            coppie.append((self.celle(rng), _criterio(self.scalare(crit))))
        n = len(coppie[0][0])
        if any(len(v) != n for v, _ in coppie):
            return ErroreExcel("#VALUE!")
        ok = [all(f(v[k]) for v, f in coppie) for k in range(n)]
        if nome == "COUNTIFS":
            return sum(ok)
        if not isinstance(bersaglio, _Intervallo):
            raise NonSupportata(f"{nome}: l'intervallo dei valori deve essere un riferimento")
        valori = self.celle(bersaglio)
        if len(valori) != n:
            return ErroreExcel("#VALUE!")
        scelti = [v for v, k in zip(valori, ok) if k and isinstance(v, (int, float)) and not isinstance(v, bool)]
        if not scelti:
            return 0
        return min(scelti) if nome == "MINIFS" else max(scelti)


RE_CELLA_FORMULA = re.compile(r'<c r="([A-Z]+\d+)"([^>]*)>(<f(?:\s[^>]*)?>[^<]*</f>|<f[^>]*/>)(?:<v\s*/>|<v></v>)</c>')


def _cella_xml(coord: str, attributi: str, formula: str, v) -> str:
    attributi = re.sub(r'\s+t="[^"]*"', "", attributi)
    if isinstance(v, ErroreExcel):
        t, testo = ' t="e"', escape(str(v))
    elif isinstance(v, bool):
        t, testo = ' t="b"', "1" if v else "0"
    elif isinstance(v, (int, float)):
        t, testo = "", (str(int(v)) if float(v).is_integer() and abs(v) < 1e15 else repr(float(v)))
    elif isinstance(v, str):
        t, testo = ' t="str"', escape(v)
    elif v is None:   # una formula che rimanda a una cella vuota vale 0 in Excel
        t, testo = "", "0"
    else:
        raise NonSupportata(f"risultato di tipo {type(v).__name__}")
    return f'<c r="{coord}"{attributi}{t}>{formula}<v>{testo}</v></c>'


def _parti_fogli(contenuti: dict[str, bytes]) -> dict[str, str]:
    """Nome del foglio → percorso della parte XML nel pacchetto."""
    from html import unescape
    wb_xml = contenuti["xl/workbook.xml"].decode("utf-8")
    rels = contenuti["xl/_rels/workbook.xml.rels"].decode("utf-8")
    target = {}
    for rel in re.findall(r"<Relationship\b[^>]*>", rels):
        rid, tgt = re.search(r'Id="([^"]+)"', rel), re.search(r'Target="([^"]+)"', rel)
        if rid and tgt:
            t = tgt.group(1)
            target[rid.group(1)] = t.lstrip("/") if t.startswith("/") else "xl/" + t
    fogli = {}
    for s in re.findall(r"<sheet\b[^>]*>", wb_xml):
        nome = re.search(r'name="([^"]*)"', s)
        rid = re.search(r'r:id="([^"]+)"', s) or re.search(r'\bid="([^"]+)"', s)
        if nome and rid and rid.group(1) in target:
            fogli[unescape(nome.group(1))] = target[rid.group(1)]
    return fogli


def riempi_cache(percorso: Path) -> list[str]:
    """Calcola le formule del file `percorso` (salvato da openpyxl) e ne scrive i risultati nella cache.
    Restituisce le celle non calcolate (funzioni non supportate), che restano senza valore in cache."""
    import openpyxl

    wb = openpyxl.load_workbook(percorso)
    calc = Calcolatore(wb)
    risultati: dict[str, dict[str, object]] = {}
    non_calcolate: list[str] = []
    for ws in wb.worksheets:
        for riga in ws.iter_rows():
            for c in riga:
                if isinstance(c.value, str) and c.value.startswith("="):
                    try:
                        risultati.setdefault(ws.title, {})[c.coordinate] = calc.valore(ws.title, c.row, c.column)
                    except NonSupportata as e:
                        non_calcolate.append(f"{ws.title}!{c.coordinate}: {e}")
    if not risultati:
        return non_calcolate

    with zipfile.ZipFile(percorso) as z:
        nomi = z.namelist()
        contenuti = {n: z.read(n) for n in nomi}
    parti = _parti_fogli(contenuti)
    for foglio, celle in risultati.items():
        parte = parti.get(foglio)
        if parte not in contenuti:
            non_calcolate.extend(f"{foglio}!{k}: parte XML non trovata" for k in celle)
            continue

        def sostituisci(m, celle=celle, foglio=foglio):
            coord, attributi, formula = m.groups()
            if coord not in celle:
                return m.group(0)
            try:
                return _cella_xml(coord, attributi, formula, celle[coord])
            except NonSupportata as e:
                non_calcolate.append(f"{foglio}!{coord}: {e}")
                return m.group(0)

        contenuti[parte] = RE_CELLA_FORMULA.sub(sostituisci, contenuti[parte].decode("utf-8")).encode("utf-8")
    with zipfile.ZipFile(percorso, "w", zipfile.ZIP_DEFLATED) as z:
        for n in nomi:
            z.writestr(n, contenuti[n])
    return non_calcolate
