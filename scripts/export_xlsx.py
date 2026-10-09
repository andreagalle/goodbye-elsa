#!/usr/bin/env python3
"""Esporta il workbook dei fondi pensione aperti (COVIP) in JSON per la dashboard.

Legge data/fondi-pensione-covip.xlsx (fonte di verità) e scrive:
  data/fondi.json, data/comparti.json, data/regole.json, data/longevita.json, data/glossario.json, data/meta.json
e ne copia una versione in docs/data/ per il sito. Rigenera anche la tabella del glossario in docs/guida/GUIDA.md
(solo il blocco tra i marcatori <!-- glossario:inizio … --> e <!-- glossario:fine -->).

openpyxl non calcola le formule: lo script risolve da solo Comparti!A (=Sheet1!$A$n) e
ricalcola le metriche H–L di Sheet1, confrontandole con i valori in cache dell'ultimo
salvataggio in Excel (warning se diversi o assenti).

Uscita con codice 1 se ci sono errori di schema. Uso:
  python scripts/export_xlsx.py [--xlsx PERCORSO] [--out DIR] [--docs DIR | --no-docs] [--check] [--strict]
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import os
import re
import subprocess
import sys
import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX_DEFAULT = ROOT / "data" / "fondi-pensione-covip.xlsx"

SHEET_FONDI = "Sheet1"
SHEET_COMPARTI = "Comparti"
SHEET_REGOLE = "Regole"        # regole generali della previdenza complementare, con fonti
SHEET_LONGEVITA = "Longevita"  # tavola di mortalità ISTAT (sopravviventi e speranza di vita)
SHEET_GLOSSARIO = "Glossario"  # termini spiegati nella dashboard (suggerimenti al passaggio del mouse) e nella guida
FONDI_PRIMA_RIGA = 3      # intestazioni alla riga 2
COMPARTI_PRIMA_RIGA = 2   # intestazioni alla riga 1

CATEGORIE = ("AZN", "BIL", "OBB MISTO", "OBB PURO", "GAR")
PERIODI = (3, 5, 10)
SI_NO = ("Sì", "No")
ONLINE = ("Sì (Ciao Elsa)", "No", "No (lista d'attesa)")

# Intestazioni attese: se il workbook cambia struttura lo script si ferma invece di esportare dati sbagliati.
INTESTAZIONI_FONDI = {
    1: "Denominazione", 2: "Pagina informativa (URL)", 3: "Società (Ciao Elsa)",
    4: "Spese di adesione (€)", 5: "Spese annue fisse (€)", 6: "Costo % sul versato",
    7: "N. linee (dichiarate Ciao Elsa)", 8: "N. comparti nel foglio Comparti",
    9: "Comm. gestione min", 10: "Comm. gestione max", 11: "Max % azioni tra i comparti",
    12: "Miglior rendimento 10 anni", 13: "Linee sostenibili (ESG)", 14: "Life cycle",
    15: "Sottoscrizione online", 16: "Scheda Ciao Elsa (fonte)", 17: "Note",
}
INTESTAZIONI_COMPARTI = {
    1: "Fondo", 2: "Comparto", 3: "Categoria (Ciao Elsa)", 4: "% Azioni", 5: "% Obbligazioni",
    6: "Rendimento netto medio annuo", 7: "Periodo rendimento (anni)",
    8: "Commissione di gestione annua", 9: "Scheda Ciao Elsa (fonte)", 10: "Note",
}

INTESTAZIONI_REGOLE = {
    1: "ID", 2: "Tema", 3: "Titolo breve", 4: "Domanda", 5: "Regola", 6: "Valore chiave",
    7: "Uguale per tutti i fondi?", 8: "Cosa varia da fondo a fondo", 9: "Riferimento normativo", 10: "Fonte",
    11: "Consultata il", 12: "In vigore dal", 13: "Note",
}
TEMI_REGOLE = ("Come funziona", "Adesione e TFR", "Tasse e deduzioni", "Prima della pensione", "Alla pensione",
               "In caso di decesso")
UGUALE_PER_TUTTI = ("Sì", "In parte", "No")
RE_ID_REGOLA = re.compile(r"^R\d{2}$")

INTESTAZIONI_LONGEVITA = {
    1: "Età", 2: "Sopravviventi – uomini", 3: "Sopravviventi – donne", 4: "Sopravviventi – uomini e donne",
    5: "Speranza di vita – uomini", 6: "Speranza di vita – donne", 7: "Speranza di vita – uomini e donne",
}
SESSI = ("uomini", "donne", "totale")
ETA_PARTENZA = 67    # età di riferimento: requisito anagrafico della pensione di vecchiaia
ETA_GRAFICO_MAX = 105

INTESTAZIONI_GLOSSARIO = {
    1: "ID", 2: "Gruppo", 3: "Termine", 4: "Per esteso", 5: "Definizione", 6: "Fonte", 7: "Consultata il", 8: "Note",
}
GRUPPI_GLOSSARIO = ("Fondi e documenti", "Investimento", "Costi e rendimenti", "Versamenti e uscite anticipate",
                    "Alla pensione", "Longevità e decesso")
RE_ID_GLOSSARIO = re.compile(r"^[a-z0-9]+(-[a-z0-9]+)*$")   # usato dalla UI: data-glossario="life-cycle"
DEFINIZIONE_MAX = 400   # oltre, il suggerimento al passaggio del mouse diventa difficile da leggere
GUIDA = ROOT / "docs" / "guida" / "GUIDA.md"
INIZIO_GLOSSARIO = "<!-- glossario:inizio"   # la tabella tra i due marcatori è generata dal foglio Glossario
FINE_GLOSSARIO = "<!-- glossario:fine -->"

# Codici delle anomalie rilevate in automatico (la descrizione finisce in meta.json per la UI).
FLAG_COMPARTI = {
    "somma_allocazione": "% azioni + % obbligazioni diversa da 100%",
    "garantito_azionario": "Comparto garantito/prudente con più del 50% di azioni",
    "categoria_incoerente": "% azioni incoerente con la categoria Ciao Elsa",
    "periodo_non_10": "Rendimento non a 10 anni: non confrontabile con gli altri",
    "rendimento_mancante": "Rendimento mancante",
    "commissione_mancante": "Commissione di gestione mancante",
}
FLAG_FONDI = {
    "linee_diverse": "N. linee dichiarate diverso dai comparti presenti nel foglio Comparti",
    "senza_dati": "Nessuna scheda né dati di dettaglio",
    "comparti_anomali": "Almeno un comparto ha anomalie: le metriche aggregate potrebbero essere falsate",
}

RE_RIF_FONDO = re.compile(r"^=\s*'?Sheet1'?!\$?A\$?(\d+)\s*$", re.IGNORECASE)
RE_GARANTITO = re.compile(r"garant|garanzia|prudent", re.IGNORECASE)
TOLLERANZA = 1e-9

# Parti generiche della denominazione COVIP, tolte per il nome breve mostrato nella UI.
RE_GENERICI = re.compile(
    r"\s*-?\s*FONDO PENSIONE APERTO\b|\bA CONTRIBUZIONE DEFINITA\b|\bPREVIDENZA COMPLEMENTARE\b"
    r"|\bIL FONDO PENSIONE APERTO DI\b|\bBY\b|\bFPA$",
    re.IGNORECASE,
)
PAROLE_MINUSCOLE = {"di", "e", "per", "il", "la", "a"}
SIGLE = {"BCC", "CNP", "BIM", "UBI", "BAP", "ZED", "ITAS", "FPA"}


class Report:
    def __init__(self) -> None:
        self.errori: list[str] = []
        self.warning: list[str] = []

    def errore(self, msg: str) -> None:
        self.errori.append(msg)

    def avviso(self, msg: str) -> None:
        self.warning.append(msg)


# ---------------------------------------------------------------- utilità

def slugify(testo: str) -> str:
    t = unicodedata.normalize("NFKD", testo).encode("ascii", "ignore").decode()
    t = t.replace("&", " e ")
    return re.sub(r"[^a-z0-9]+", "-", t.lower()).strip("-")


def nome_breve(denominazione: str) -> str:
    base = RE_GENERICI.sub(" ", denominazione)
    base = re.sub(r"\s+", " ", base).strip(" -")
    parole = []
    for i, p in enumerate(base.split(" ")):
        if p.upper() in SIGLE:
            parole.append(p.upper())
        elif i > 0 and p.lower() in PAROLE_MINUSCOLE:
            parole.append(p.lower())
        else:
            parole.append(p.capitalize())
    return " ".join(parole) or denominazione


def testo(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def numero(v, campo: str, dove: str, rep: Report) -> float | None:
    """Celle vuote → None (mai 0). Testo non numerico → errore."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, bool):
        rep.errore(f"{dove}: {campo} è un booleano ({v!r})")
        return None
    if isinstance(v, (int, float)):
        return float(v)
    rep.errore(f"{dove}: {campo} non numerico ({v!r})")
    return None


def arrotonda(v: float | None, cifre: int = 6) -> float | None:
    """Toglie il rumore binario (0.006500000000000001 → 0.0065)."""
    if v is None:
        return None
    r = round(v, cifre)
    return int(r) if r == int(r) and abs(r) >= 1 else r


def link(cell) -> str | None:
    if cell.hyperlink is not None and cell.hyperlink.target:
        return cell.hyperlink.target.strip()
    v = testo(cell.value)
    if v and re.match(r"^https?://", v):
        return v
    return None


def uguale(a, b) -> bool:
    if a is None or b is None:
        return a is None and b is None
    return abs(float(a) - float(b)) <= TOLLERANZA


def versione_app() -> str:
    """Versione pubblicata: APP_VERSION in CI (calcolata da scripts/versione.py), altrimenti git describe."""
    v = os.environ.get("APP_VERSION")
    if v:
        return v
    try:
        return subprocess.run(
            ["git", "describe", "--tags", "--match", "v*", "--always", "--dirty"],
            cwd=ROOT, capture_output=True, text=True, check=True,
        ).stdout.strip() + " (sviluppo)"
    except (OSError, subprocess.CalledProcessError):
        return "sviluppo"


# Riconoscimento della licenza dal file LICENSE del repository (la UI non la scrive mai a mano).
LICENZE = [
    ("unencumbered software released into the public domain", "The Unlicense", "Unlicense"),
    ("mit license", "MIT License", "MIT"),
    ("apache license", "Apache License 2.0", "Apache-2.0"),
    ("gnu affero general public license", "GNU AGPL v3", "AGPL-3.0"),
    ("gnu lesser general public license", "GNU LGPL", "LGPL"),
    ("gnu general public license", "GNU GPL", "GPL"),
    ("mozilla public license", "Mozilla Public License 2.0", "MPL-2.0"),
    ("creative commons", "Creative Commons", "CC"),
]


def licenza_repo(rep: "Report") -> dict | None:
    f = ROOT / "LICENSE"
    if not f.exists():
        rep.avviso("File LICENSE assente: la licenza non viene mostrata nel sito")
        return None
    testo_lic = " ".join(f.read_text(encoding="utf-8", errors="replace").lower().split())
    for chiave, nome, spdx in LICENZE:
        if chiave in testo_lic:
            return {"nome": nome, "spdx": spdx, "file": "LICENSE"}
    rep.avviso("Licenza nel file LICENSE non riconosciuta: aggiungerla a LICENZE in export_xlsx.py")
    return {"nome": "vedi file LICENSE", "spdx": None, "file": "LICENSE"}


def git_commit() -> str | None:
    sha = os.environ.get("GITHUB_SHA")
    if sha:
        return sha
    try:
        return subprocess.run(
            ["git", "rev-parse", "HEAD"], cwd=ROOT, capture_output=True, text=True, check=True
        ).stdout.strip() or None
    except (OSError, subprocess.CalledProcessError):
        return None


# ---------------------------------------------------------------- lettura

def verifica_intestazioni(ws, riga: int, attese: dict[int, str], rep: Report) -> None:
    for col, nome in attese.items():
        trovato = testo(ws.cell(riga, col).value)
        if trovato != nome:
            rep.errore(f"{ws.title}!{ws.cell(riga, col).coordinate}: intestazione {trovato!r}, attesa {nome!r}")


def leggi_fondi(ws, ws_cache, rep: Report) -> list[dict]:
    fondi = []
    for r in range(FONDI_PRIMA_RIGA, ws.max_row + 1):
        den = testo(ws.cell(r, 1).value)
        if den is None:
            if any(ws.cell(r, c).value is not None for c in range(2, 18)):
                rep.errore(f"Sheet1!A{r}: denominazione vuota in una riga con dati")
            continue
        dove = f"Sheet1 riga {r}"
        cache = {c: ws_cache.cell(r, c).value for c in range(8, 13)}
        n_linee = numero(ws.cell(r, 7).value, "N. linee", dove, rep)
        fondi.append({
            "id": None,  # assegnato dopo, garantendo l'unicità
            "riga": r,
            "denominazione": den,
            "nome_breve": nome_breve(den),
            "url": link(ws.cell(r, 2)),
            "societa": testo(ws.cell(r, 3).value),
            "spese_adesione": numero(ws.cell(r, 4).value, "Spese di adesione", dove, rep),
            "spese_annue": numero(ws.cell(r, 5).value, "Spese annue fisse", dove, rep),
            "costo_pct_versato": numero(ws.cell(r, 6).value, "Costo % sul versato", dove, rep),
            "n_linee": int(n_linee) if n_linee is not None else None,
            "esg": testo(ws.cell(r, 13).value),
            "life_cycle": testo(ws.cell(r, 14).value),
            "online": testo(ws.cell(r, 15).value),
            "scheda_url": link(ws.cell(r, 16)),
            "note": testo(ws.cell(r, 17).value),
            "_cache": cache,
        })
    return fondi


def leggi_comparti(ws, fondi_per_riga: dict[int, dict], fondi_per_nome: dict[str, dict], rep: Report) -> list[dict]:
    comparti = []
    for r in range(COMPARTI_PRIMA_RIGA, ws.max_row + 1):
        if all(ws.cell(r, c).value is None for c in range(1, 11)):
            continue
        dove = f"Comparti riga {r}"
        raw = ws.cell(r, 1).value
        fondo = None
        if isinstance(raw, str) and raw.startswith("="):
            m = RE_RIF_FONDO.match(raw)
            if not m:
                rep.errore(f"{dove}: formula in colonna A non riconosciuta ({raw!r}), attesa =Sheet1!$A$n")
            else:
                fondo = fondi_per_riga.get(int(m.group(1)))
                if fondo is None:
                    rep.errore(f"{dove}: {raw} punta a una riga di Sheet1 senza fondo")
        else:
            nome = testo(raw)
            fondo = fondi_per_nome.get(nome or "")
            if fondo is None:
                rep.errore(f"{dove}: fondo {nome!r} non trovato in Sheet1")
            else:
                rep.avviso(f"{dove}: colonna A contiene testo invece della formula =Sheet1!$A${fondo['riga']}")
        if fondo is None:
            continue
        periodo = numero(ws.cell(r, 7).value, "Periodo rendimento", dove, rep)
        comparti.append({
            "fondo_id": fondo["id"],
            "riga": r,
            "comparto": testo(ws.cell(r, 2).value),
            "categoria": testo(ws.cell(r, 3).value),
            "azioni": numero(ws.cell(r, 4).value, "% Azioni", dove, rep),
            "obbligazioni": numero(ws.cell(r, 5).value, "% Obbligazioni", dove, rep),
            "rendimento": numero(ws.cell(r, 6).value, "Rendimento", dove, rep),
            "periodo_anni": int(periodo) if periodo is not None else None,
            "commissione": numero(ws.cell(r, 8).value, "Commissione", dove, rep),
            "scheda_url": link(ws.cell(r, 9)),
            "note": testo(ws.cell(r, 10).value),
            "flag_anomalia": [],
        })
    return comparti


def data_iso(v) -> str | None:
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    return None


def leggi_regole(ws, rep: Report) -> list[dict]:
    regole = []
    for r in range(2, ws.max_row + 1):
        if all(ws.cell(r, c).value is None for c in range(1, 14)):
            continue
        dove = f"Regole riga {r}"
        fonte = ws.cell(r, 10)
        regola = {
            "id": testo(ws.cell(r, 1).value),
            "tema": testo(ws.cell(r, 2).value),
            "titolo": testo(ws.cell(r, 3).value),
            "domanda": testo(ws.cell(r, 4).value),
            "regola": testo(ws.cell(r, 5).value),
            "valore": testo(ws.cell(r, 6).value),
            "uguale_per_tutti": testo(ws.cell(r, 7).value),
            "varia": testo(ws.cell(r, 8).value),
            "riferimento": testo(ws.cell(r, 9).value),
            "fonte_nome": testo(fonte.value),
            "fonte_url": link(fonte),
            "consultata_il": data_iso(ws.cell(r, 11).value),
            "in_vigore_dal": data_iso(ws.cell(r, 12).value),
            "note": testo(ws.cell(r, 13).value),
        }
        for c, campo in ((11, "consultata_il"), (12, "in_vigore_dal")):
            v = ws.cell(r, c).value
            if v not in (None, "") and regola[campo] is None:
                rep.errore(f"{dove}: {campo} non è una data ({v!r})")
        regola["_dove"] = dove
        regole.append(regola)
    return regole


def valida_regole(regole: list[dict], rep: Report) -> None:
    visti = set()
    for g in regole:
        dove = g["_dove"]
        if not g["id"] or not RE_ID_REGOLA.match(g["id"]):
            rep.errore(f"{dove}: ID {g['id']!r} non valido (atteso R01, R02, …)")
        elif g["id"] in visti:
            rep.errore(f"{dove}: ID duplicato {g['id']}")
        visti.add(g["id"])
        if g["tema"] not in TEMI_REGOLE:
            rep.errore(f"{dove}: tema {g['tema']!r} non in {TEMI_REGOLE}")
        for campo in ("titolo", "domanda", "regola", "riferimento"):
            if not g[campo]:
                rep.errore(f"{dove}: {campo} vuoto")
        if g["uguale_per_tutti"] not in UGUALE_PER_TUTTI:
            rep.errore(f"{dove}: 'Uguale per tutti i fondi?' = {g['uguale_per_tutti']!r}, ammessi {UGUALE_PER_TUTTI}")
        elif g["uguale_per_tutti"] != "Sì" and not g["varia"]:
            rep.errore(f"{dove}: indicare cosa varia da fondo a fondo")
        if not (g["fonte_nome"] and g["fonte_url"]):
            rep.errore(f"{dove}: fonte senza nome o senza link")
        if not g["consultata_il"]:
            rep.errore(f"{dove}: manca la data di consultazione della fonte")


def leggi_longevita(ws, rep: Report) -> dict:
    """Tavola ISTAT: sopravviventi e speranza di vita per età; le % di sopravvivenza si ricalcolano qui."""
    righe = []
    for r in range(2, ws.max_row + 1):
        eta = ws.cell(r, 1).value
        if eta is None:
            continue
        dove = f"Longevita riga {r}"
        if not isinstance(eta, int):
            rep.errore(f"{dove}: età non intera ({eta!r})")
            continue
        riga = {"eta": eta}
        for j, sesso in enumerate(SESSI):
            riga[f"l_{sesso}"] = numero(ws.cell(r, 2 + j).value, f"sopravviventi {sesso}", dove, rep)
            riga[f"e_{sesso}"] = numero(ws.cell(r, 5 + j).value, f"speranza di vita {sesso}", dove, rep)
        righe.append(riga)
    meta = {testo(ws.cell(r, 12).value): ws.cell(r, 13) for r in range(1, 6) if testo(ws.cell(r, 12).value)}
    fonte = {
        "nome": testo(meta["Fonte"].value) if "Fonte" in meta else None,
        "url": link(meta["Link"]) if "Link" in meta else None,
        "consultata_il": data_iso(meta["Consultata il"].value) if "Consultata il" in meta else None,
        "note": testo(meta["Note"].value) if "Note" in meta else None,
    }
    if not (fonte["nome"] and fonte["url"] and fonte["consultata_il"]):
        rep.errore("Longevita: mancano fonte, link o data di consultazione (celle L1:M3)")

    eta = [x["eta"] for x in righe]
    if eta != list(range(eta[0], eta[0] + len(eta))) if eta else True:
        rep.errore("Longevita: le età devono essere consecutive")
        return {}
    per_eta = {x["eta"]: x for x in righe}
    if ETA_PARTENZA not in per_eta:
        rep.errore(f"Longevita: manca l'età {ETA_PARTENZA}")
        return {}
    for sesso in SESSI:
        valori = [per_eta[e][f"l_{sesso}"] for e in eta]
        if any(v is None or v < 0 for v in valori) or any(b > a for a, b in zip(valori, valori[1:])):
            rep.errore(f"Longevita: sopravviventi {sesso} mancanti, negativi o crescenti con l'età")
            return {}

    base = {s: per_eta[ETA_PARTENZA][f"l_{s}"] for s in SESSI}
    serie, sintesi = {}, {}
    for s in SESSI:
        quota = {e: per_eta[e][f"l_{s}"] / base[s] for e in eta if ETA_PARTENZA <= e <= ETA_GRAFICO_MAX}
        serie[s] = [{"eta": e, "vivi": round(q, 4)} for e, q in quota.items()]

        def eta_soglia(soglia):
            return next((e for e, q in quota.items() if q <= soglia), None)

        e67 = per_eta[ETA_PARTENZA][f"e_{s}"]
        sintesi[s] = {
            "speranza": round(e67, 2),
            "eta_75_vivi": eta_soglia(0.75),
            "eta_50_vivi": eta_soglia(0.5),
            "eta_25_vivi": eta_soglia(0.25),
            "eta_10_vivi": eta_soglia(0.10),
        }
    # rendita a durata definita (art. 11 c. 3-ter): anni interi di speranza di vita, tavola uomini e donne insieme
    durata = int(sintesi["totale"]["speranza"])
    fine = ETA_PARTENZA + durata
    durata_definita = {
        "eta_inizio": ETA_PARTENZA,
        "anni": durata,
        "eta_fine": fine,
        "vivi_a_fine": {s: round(per_eta[fine][f"l_{s}"] / base[s], 4) for s in SESSI},
    }
    return {"fonte": fonte, "eta_partenza": ETA_PARTENZA, "sintesi": sintesi,
            "durata_definita": durata_definita, "serie": serie}


def leggi_glossario(ws, rep: Report) -> list[dict]:
    voci = []
    for r in range(2, ws.max_row + 1):
        if all(ws.cell(r, c).value is None for c in range(1, 9)):
            continue
        fonte = ws.cell(r, 6)
        voce = {
            "id": testo(ws.cell(r, 1).value),
            "gruppo": testo(ws.cell(r, 2).value),
            "termine": testo(ws.cell(r, 3).value),
            "esteso": testo(ws.cell(r, 4).value),
            "definizione": testo(ws.cell(r, 5).value),
            "fonte_nome": testo(fonte.value),
            "fonte_url": link(fonte),
            "consultata_il": data_iso(ws.cell(r, 7).value),
            "note": testo(ws.cell(r, 8).value),
            "_dove": f"Glossario riga {r}",
        }
        v = ws.cell(r, 7).value
        if v not in (None, "") and voce["consultata_il"] is None:
            rep.errore(f"{voce['_dove']}: consultata_il non è una data ({v!r})")
        voci.append(voce)
    return voci


def valida_glossario(voci: list[dict], rep: Report) -> None:
    ids, termini = set(), set()
    for v in voci:
        dove = v["_dove"]
        if not v["id"] or not RE_ID_GLOSSARIO.match(v["id"]):
            rep.errore(f"{dove}: ID {v['id']!r} non valido (minuscole, cifre e trattini, es. life-cycle)")
        elif v["id"] in ids:
            rep.errore(f"{dove}: ID duplicato {v['id']}")
        ids.add(v["id"])
        if v["gruppo"] not in GRUPPI_GLOSSARIO:
            rep.errore(f"{dove}: gruppo {v['gruppo']!r} non in {GRUPPI_GLOSSARIO}")
        for campo in ("termine", "definizione"):
            if not v[campo]:
                rep.errore(f"{dove}: {campo} vuoto")
        if v["termine"] and v["termine"].lower() in termini:
            rep.errore(f"{dove}: termine duplicato {v['termine']!r}")
        termini.add((v["termine"] or "").lower())
        if not (v["fonte_nome"] and v["fonte_url"]):
            rep.errore(f"{dove}: fonte senza nome o senza link")
        if not v["consultata_il"]:
            rep.errore(f"{dove}: manca la data di consultazione della fonte")
        if v["definizione"] and len(v["definizione"]) > DEFINIZIONE_MAX:
            rep.avviso(f"{dove}: definizione di {len(v['definizione'])} caratteri, troppo lunga per un suggerimento "
                       f"(massimo {DEFINIZIONE_MAX}): il resto può andare in Note")


def fonte_breve(nome: str) -> str:
    """Nome corto per i link compatti: la parte prima del trattino ("COVIP – Glossario" → "COVIP")."""
    return nome.split(" – ")[0].strip()


def url_comune(urls: list[str]) -> str:
    """Un solo link per una fonte citata con più link (le lettere del glossario COVIP → la pagina del glossario)."""
    pagine = sorted({u.split("#")[0] for u in urls})
    if len(pagine) == 1:
        return pagine[0]
    comune = os.path.commonprefix(pagine)
    return comune[: comune.rfind("/")]


def glossario_md(voci: list[dict]) -> str:
    """Il blocco del glossario per docs/guida/GUIDA.md: una tabella per gruppo, nell'ordine del foglio."""
    def cella(s: str) -> str:
        return s.replace("|", "\\|").replace("\n", " ")

    righe = [f"{INIZIO_GLOSSARIO} — generato da scripts/export_xlsx.py dal foglio Glossario: non modificare a mano -->"]
    for gruppo in dict.fromkeys(v["gruppo"] for v in voci):
        righe += ["", f"### {gruppo}", "", "| Termine | Significato | Fonte |", "|---|---|---|"]
        for v in (x for x in voci if x["gruppo"] == gruppo):
            termine = f"**{v['termine']}**" + (f" ({v['esteso']})" if v["esteso"] else "")
            spiegazione = v["definizione"] + (f"<br>*{v['note'].replace('*', '')}*" if v["note"] else "")
            fonte = f"[{fonte_breve(v['fonte_nome'])}]({v['fonte_url']})"
            righe.append(f"| {cella(termine)} | {cella(spiegazione)} | {cella(fonte)} |")
    righe += ["", FINE_GLOSSARIO]
    return "\n".join(righe)


def aggiorna_guida(voci: list[dict]) -> str | None:
    """Riscrive in GUIDA.md solo il blocco tra i marcatori del glossario. Restituisce un errore, o None."""
    testo_guida = GUIDA.read_text(encoding="utf-8")
    i, j = testo_guida.find(INIZIO_GLOSSARIO), testo_guida.find(FINE_GLOSSARIO)
    if i < 0 or j < i:
        return f"{GUIDA.relative_to(ROOT)}: mancano i marcatori {INIZIO_GLOSSARIO} … --> e {FINE_GLOSSARIO}"
    nuovo = testo_guida[:i] + glossario_md(voci) + testo_guida[j + len(FINE_GLOSSARIO):]
    if nuovo != testo_guida:
        GUIDA.write_text(nuovo, encoding="utf-8")
        print(f"Aggiornata la tabella del glossario in {GUIDA.relative_to(ROOT)}")
    return None


# ---------------------------------------------------------------- calcoli

def ricalcola_metriche(fondo: dict, suoi: list[dict]) -> None:
    """Equivalente Python delle formule H–L di Sheet1."""
    def valori(campo, filtro=lambda c: True):
        return [c[campo] for c in suoi if c[campo] is not None and filtro(c)]

    fondo["n_comparti"] = len(suoi)
    comm = valori("commissione")
    azioni = valori("azioni")
    rend10 = valori("rendimento", lambda c: c["periodo_anni"] == 10)
    vuoto = len(suoi) == 0
    fondo["comm_min"] = None if vuoto or not comm else min(comm)
    fondo["comm_max"] = None if vuoto or not comm else max(comm)
    fondo["max_azioni"] = None if vuoto or not azioni else max(azioni)
    fondo["best_rend_10a"] = max(rend10) if rend10 else None


def confronta_cache(fondo: dict, rep: Report, cache_vuote: list[str]) -> None:
    campi = {8: "n_comparti", 9: "comm_min", 10: "comm_max", 11: "max_azioni", 12: "best_rend_10a"}
    lettere = {8: "H", 9: "I", 10: "J", 11: "K", 12: "L"}
    for col, campo in campi.items():
        cella = f"Sheet1!{lettere[col]}{fondo['riga']}"
        in_cache = fondo["_cache"][col]
        calcolato = fondo[campo]
        if in_cache is None and fondo["n_comparti"] > 0:
            cache_vuote.append(cella)
            continue
        if in_cache == "":
            in_cache = None
        if isinstance(in_cache, str):
            rep.avviso(f"{cella}: valore in cache non numerico ({in_cache!r})")
            continue
        if col == 8 and in_cache is None:
            in_cache = 0 if fondo["n_comparti"] == 0 else None
        if not uguale(in_cache, calcolato):
            rep.avviso(f"{cella} ({campo}): in cache {in_cache!r}, ricalcolato {calcolato!r}")


def flag_comparto(c: dict) -> list[str]:
    flag = []
    a, o = c["azioni"], c["obbligazioni"]
    if a is not None and o is not None and abs(a + o - 1) > 0.005:
        flag.append("somma_allocazione")
    if a is not None and a > 0.5 and (RE_GARANTITO.search(c["comparto"] or "") or c["categoria"] == "GAR"):
        flag.append("garantito_azionario")
    if a is not None and "garantito_azionario" not in flag and (
        (c["categoria"] == "AZN" and a < 0.5)
        or (c["categoria"] in ("OBB PURO", "OBB MISTO", "GAR") and a > 0.5)
    ):
        flag.append("categoria_incoerente")
    if c["periodo_anni"] is not None and c["periodo_anni"] != 10:
        flag.append("periodo_non_10")
    if c["rendimento"] is None:
        flag.append("rendimento_mancante")
    if c["commissione"] is None:
        flag.append("commissione_mancante")
    return flag


# ---------------------------------------------------------------- validazione

def valida(fondi: list[dict], comparti: list[dict], rep: Report) -> None:
    ids = set()
    for f in fondi:
        dove = f"Sheet1 riga {f['riga']} ({f['nome_breve']})"
        if f["id"] in ids:
            rep.errore(f"{dove}: id duplicato {f['id']}")
        ids.add(f["id"])
        if not f["url"]:
            rep.errore(f"{dove}: manca la pagina informativa (colonna B)")
        for campo, ammessi in (("esg", SI_NO), ("life_cycle", SI_NO), ("online", ONLINE)):
            if f[campo] is not None and f[campo] not in ammessi:
                rep.errore(f"{dove}: {campo} = {f[campo]!r}, ammessi {ammessi}")
        for campo in ("spese_adesione", "spese_annue"):
            if f[campo] is not None and f[campo] < 0:
                rep.errore(f"{dove}: {campo} negativo")
        if f["costo_pct_versato"] is not None and not 0 <= f["costo_pct_versato"] < 1:
            rep.errore(f"{dove}: costo_pct_versato fuori da [0, 1)")
        if f["has_dati"] and not (f["scheda_url"] or f["note"]):
            rep.errore(f"{dove}: dati senza fonte (né scheda né nota)")

    for c in comparti:
        dove = f"Comparti riga {c['riga']}"
        if c["fondo_id"] not in ids:
            rep.errore(f"{dove}: fondo_id {c['fondo_id']} inesistente")
        if not c["comparto"]:
            rep.errore(f"{dove}: nome del comparto vuoto")
        if c["categoria"] not in CATEGORIE:
            rep.errore(f"{dove}: categoria {c['categoria']!r} non in {CATEGORIE}")
        for campo in ("azioni", "obbligazioni", "commissione"):
            v = c[campo]
            if v is not None and not 0 <= v <= 1:
                rep.errore(f"{dove}: {campo} = {v} fuori da [0, 1] (le percentuali sono decimali)")
        if c["commissione"] is not None and c["commissione"] > 0.05:
            rep.avviso(f"{dove}: commissione {c['commissione']:.2%} insolitamente alta")
        if c["rendimento"] is not None and not -0.5 <= c["rendimento"] <= 0.5:
            rep.errore(f"{dove}: rendimento {c['rendimento']} fuori scala (decimale atteso)")
        if c["periodo_anni"] not in (*PERIODI, None):
            rep.errore(f"{dove}: periodo_anni {c['periodo_anni']} non in {PERIODI}")
        if c["rendimento"] is not None and c["periodo_anni"] is None:
            rep.errore(f"{dove}: rendimento senza periodo")
        if not (c["scheda_url"] or c["note"]):
            rep.errore(f"{dove}: comparto senza fonte (né scheda né nota)")


# ---------------------------------------------------------------- main

def esporta(xlsx: Path, rep: Report) -> tuple[list[dict], list[dict], dict]:
    """Restituisce (fondi, comparti, meta); regole, longevità e glossario finiscono in meta["_extra"] e vengono
    scritti in file separati da main()."""
    wb = openpyxl.load_workbook(xlsx)
    wb_cache = openpyxl.load_workbook(xlsx, data_only=True)
    for nome in (SHEET_FONDI, SHEET_COMPARTI, SHEET_REGOLE, SHEET_LONGEVITA, SHEET_GLOSSARIO):
        if nome not in wb.sheetnames:
            rep.errore(f"Foglio {nome!r} mancante (presenti: {wb.sheetnames})")
    if rep.errori:
        return [], [], {}

    ws_f, ws_c = wb[SHEET_FONDI], wb[SHEET_COMPARTI]
    verifica_intestazioni(ws_f, 2, INTESTAZIONI_FONDI, rep)
    verifica_intestazioni(ws_c, 1, INTESTAZIONI_COMPARTI, rep)
    verifica_intestazioni(wb[SHEET_REGOLE], 1, INTESTAZIONI_REGOLE, rep)
    verifica_intestazioni(wb[SHEET_LONGEVITA], 1, INTESTAZIONI_LONGEVITA, rep)
    verifica_intestazioni(wb[SHEET_GLOSSARIO], 1, INTESTAZIONI_GLOSSARIO, rep)
    if rep.errori:
        return [], [], {}

    regole = leggi_regole(wb[SHEET_REGOLE], rep)
    valida_regole(regole, rep)
    regole_out = [{k: v for k, v in g.items() if k != "_dove"} for g in regole]
    longevita = leggi_longevita(wb[SHEET_LONGEVITA], rep)
    glossario = leggi_glossario(wb[SHEET_GLOSSARIO], rep)
    valida_glossario(glossario, rep)
    glossario_out = [{k: v for k, v in g.items() if k != "_dove"} for g in glossario]

    fondi = leggi_fondi(ws_f, wb_cache[SHEET_FONDI], rep)
    usati: set[str] = set()
    for f in fondi:
        base = slugify(f["nome_breve"]) or f"fondo-{f['riga']}"
        sid, n = base, 2
        while sid in usati:
            sid, n = f"{base}-{n}", n + 1
        usati.add(sid)
        f["id"] = sid

    per_riga = {f["riga"]: f for f in fondi}
    per_nome = {f["denominazione"]: f for f in fondi}
    comparti = leggi_comparti(ws_c, per_riga, per_nome, rep)

    cache_vuote: list[str] = []
    for f in fondi:
        suoi = [c for c in comparti if c["fondo_id"] == f["id"]]
        ricalcola_metriche(f, suoi)
        confronta_cache(f, rep, cache_vuote)
    if cache_vuote:
        rep.avviso(
            f"{len(cache_vuote)} celle con formula senza valore in cache (file salvato fuori da Excel?): "
            "uso i valori ricalcolati in Python"
        )

    for c in comparti:
        c["flag_anomalia"] = flag_comparto(c)

    for f in fondi:
        suoi = [c for c in comparti if c["fondo_id"] == f["id"]]
        f["has_dati"] = f["n_comparti"] > 0 or f["societa"] is not None
        flag = []
        if not f["has_dati"]:
            flag.append("senza_dati")
        elif f["n_linee"] is not None and f["n_linee"] != f["n_comparti"]:
            flag.append("linee_diverse")
        if any(set(c["flag_anomalia"]) & {"somma_allocazione", "garantito_azionario", "categoria_incoerente"}
               for c in suoi):
            flag.append("comparti_anomali")
        f["flag_anomalia"] = flag

    valida(fondi, comparti, rep)

    ordine_fondi = [
        "id", "riga", "denominazione", "nome_breve", "url", "societa", "spese_adesione", "spese_annue",
        "costo_pct_versato", "n_linee", "n_comparti", "comm_min", "comm_max", "max_azioni", "best_rend_10a",
        "esg", "life_cycle", "online", "scheda_url", "note", "has_dati", "flag_anomalia",
    ]
    ordine_comparti = [
        "fondo_id", "riga", "comparto", "categoria", "azioni", "obbligazioni", "rendimento", "periodo_anni",
        "commissione", "scheda_url", "note", "flag_anomalia",
    ]
    numerici = {"spese_adesione", "spese_annue", "costo_pct_versato", "comm_min", "comm_max", "max_azioni",
                "best_rend_10a", "azioni", "obbligazioni", "rendimento", "commissione"}

    def pulisci(d: dict, ordine: list[str]) -> dict:
        return {k: (arrotonda(d[k]) if k in numerici else d[k]) for k in ordine}

    fondi_out = [pulisci(f, ordine_fondi) for f in fondi]
    comparti_out = [pulisci(c, ordine_comparti) for c in comparti]

    props = wb.properties
    modificato = props.modified.isoformat() if props.modified else None
    meta = {
        "versione": versione_app(),
        "generato_il": dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat(),
        "workbook": xlsx.name,
        "workbook_modificato_il": modificato,
        "workbook_sha256": hashlib.sha256(xlsx.read_bytes()).hexdigest(),
        "commit": git_commit(),
        "licenza": licenza_repo(rep),
        "conteggi": {
            "fondi": len(fondi_out),
            "fondi_con_dati": sum(f["has_dati"] for f in fondi_out),
            "comparti": len(comparti_out),
            "comparti_10_anni": sum(c["periodo_anni"] == 10 for c in comparti_out),
            "comparti_con_anomalie": sum(bool(c["flag_anomalia"]) for c in comparti_out),
            "fondi_con_anomalie": sum(bool(f["flag_anomalia"]) for f in fondi_out),
            "regole": len(regole_out),
            "glossario": len(glossario_out),
        },
        "flag": {"comparti": FLAG_COMPARTI, "fondi": FLAG_FONDI},
        "avvisi_export": rep.warning,
        "fonti": [
            {"nome": "COVIP – Elenco dei fondi iscritti all'Albo (Sezione II – Fondi pensione aperti)",
             "url": "https://www.covip.it/la-covip-e-la-sua-attivita/albo-fondi-pensione/elenco-fondi-albo", "tipo": "primaria",
             "dettaglio": "filtro Tipologia = Sezione II – Fondi pensione aperti"},
            {"nome": "Pagine informative dei gestori (Nota informativa, Scheda costi)", "url": None,
             "tipo": "primaria", "dettaglio": "colonna url di fondi.json"},
            {"nome": "Generali – confluenza di Almeglio in Generali Global dal 1/1/2027",
             "url": "https://www.generali.it", "tipo": "primaria"},
            {"nome": "Ciao Elsa – schede dei fondi pensione aperti",
             "url": "https://www.ciaoelsa.com/schede-fondo/fondi-pensione-aperti/", "tipo": "secondaria",
             "dettaglio": "colonna scheda_url di fondi.json e comparti.json"},
        ],
    }
    # fonti delle regole e della tavola di longevità, senza duplicati (la lista resta allineata all'Excel)
    viste = {f["url"] for f in meta["fonti"]}
    for g in regole_out:
        if g["fonte_url"] not in viste:
            viste.add(g["fonte_url"])
            meta["fonti"].append({"nome": g["fonte_nome"], "url": g["fonte_url"], "tipo": "primaria",
                                  "dettaglio": "foglio Regole"})
    if longevita.get("fonte", {}).get("url") and longevita["fonte"]["url"] not in viste:
        meta["fonti"].append({"nome": longevita["fonte"]["nome"], "url": longevita["fonte"]["url"],
                              "tipo": "primaria", "dettaglio": "foglio Longevita"})
        viste.add(longevita["fonte"]["url"])
    # fonti del glossario: una sola voce per fonte (le voci del glossario COVIP puntano a lettere diverse)
    link_per_fonte: dict[str, list[str]] = {}
    for g in glossario_out:
        if g["fonte_nome"] and g["fonte_url"]:
            link_per_fonte.setdefault(g["fonte_nome"], []).append(g["fonte_url"])
    for nome, urls in link_per_fonte.items():
        url = url_comune(urls)
        if url not in viste:
            viste.add(url)
            meta["fonti"].append({"nome": nome, "url": url, "dettaglio": "foglio Glossario",
                                  "tipo": "secondaria" if nome.lower().startswith("ciao elsa") else "primaria"})
    meta["_extra"] = {"regole": regole_out, "longevita": longevita, "glossario": glossario_out}
    return fondi_out, comparti_out, meta


def scrivi_json(percorso: Path, dati) -> None:
    percorso.parent.mkdir(parents=True, exist_ok=True)
    percorso.write_text(json.dumps(dati, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--xlsx", type=Path, default=XLSX_DEFAULT)
    ap.add_argument("--out", type=Path, default=ROOT / "data", help="cartella dei JSON generati")
    ap.add_argument("--docs", type=Path, default=ROOT / "docs" / "data",
                    help="cartella del sito in cui copiare i JSON")
    ap.add_argument("--no-docs", action="store_true", help="non copiare i JSON nel sito")
    ap.add_argument("--check", action="store_true", help="valida soltanto, senza scrivere file")
    ap.add_argument("--strict", action="store_true", help="tratta anche i warning come errori")
    args = ap.parse_args(argv)

    rep = Report()
    fondi, comparti, meta = esporta(args.xlsx, rep)

    for w in rep.warning:
        print(f"WARNING: {w}", file=sys.stderr)
    for e in rep.errori:
        print(f"ERRORE:  {e}", file=sys.stderr)
    if rep.errori or (args.strict and rep.warning):
        print(f"Export fallito: {len(rep.errori)} errori, {len(rep.warning)} warning.", file=sys.stderr)
        return 1

    extra = meta.pop("_extra")
    c = meta["conteggi"]
    print(f"OK: {c['fondi']} fondi ({c['fondi_con_dati']} con dati), {c['comparti']} comparti, "
          f"{c['comparti_con_anomalie']} comparti segnalati, {c['regole']} regole, {c['glossario']} voci di glossario, "
          f"{len(rep.warning)} warning.")
    if args.check:
        return 0

    for cartella in [args.out] + ([] if args.no_docs else [args.docs]):
        scrivi_json(cartella / "fondi.json", fondi)
        scrivi_json(cartella / "comparti.json", comparti)
        scrivi_json(cartella / "regole.json", extra["regole"])
        scrivi_json(cartella / "longevita.json", extra["longevita"])
        scrivi_json(cartella / "glossario.json", extra["glossario"])
        scrivi_json(cartella / "meta.json", meta)
        print(f"Scritto in {cartella}")
    if not args.no_docs:
        errore = aggiorna_guida(extra["glossario"])
        if errore:
            print(f"ERRORE:  {errore}", file=sys.stderr)
            return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
